"""The final result: one outcome and one reason per paper, combining every step in the order the human team works.

    1. Titles          the panel (excluded only if all 3 exclude)
    2. Abstracts       the panel (weighted majority), then the age rule (65+) on the abstract: excluded only when
                       the abstract clearly shows participants under 65 (age_check_abstract.csv)
    3. Full text       free full text only. The panel (majority), then the second look on papers it left unclear
                       (second_look.csv), then the age rule (65+) on the full text (age_check.csv)
    4. Outcome         included / excluded at <stage> / for the professor / not assessed (no free full text)

Inputs (all in the run folder): decisions.jsonl + in_progress.jsonl, second_look.csv, age_check.csv,
age_check_abstract.csv, and free_fulltext_found.csv (from scripts/fetch_free_fulltext.py; may be in an earlier
run folder: --free-found) to tell papers with a free copy still to save by hand from paid-access ones.

    python -m adrd_review.final_results data/runs/PROF_search_free/<timestamp> [--free-found <csv>]

Writes <run>/final/: decisions.csv (one row per paper), summary.json, stage_table.md, for_professor.csv,
not_assessed.csv, the final charts (final_stages.png, final_reasons.png + their CSVs), and the professor's blind
random check (professor_check.csv / .xlsx: 10 included + 10 excluded at each stage, shuffled; answer key separate).
"""

from __future__ import annotations

import argparse
import csv
import json
import sys
from collections import Counter
from pathlib import Path

from adrd_review.codebook import load_codebook
from adrd_review.eval.stages import collect

PAID, HAND = "not assessed: paid access", "not assessed: free copy to save by hand"


def _read(path: Path) -> dict[str, dict]:
    return {r["pmid"]: r for r in csv.DictReader(path.open())} if path.exists() else {}


def _failing_rule(marks: str, names: dict[str, str]) -> str:
    """First rule a second-look answer marks NOT MET / APPLIES (its rule-mark string), by codebook order."""
    bad = [p.split(" ")[0] for p in marks.split(" · ") if p.endswith(("not_met", "NOT MET", "applies", "APPLIES"))]
    order = list(names)
    bad.sort(key=lambda r: order.index(r) if r in order else 99)
    return names.get(bad[0], bad[0]) if bad else "rule clearly not met"


def _open_rules(marks: str) -> list[str]:
    return [p.split(" ")[0] for p in marks.split(" · ") if p.endswith("unclear") and p.startswith("I")]


def build(run_dir: Path, free_found: Path | None = None) -> list[dict]:
    meta = json.loads((run_dir / "run_meta.json").read_text())
    codebook = load_codebook(meta["config"]["codebook"])
    i4 = "I4 Older adults (65 or over)"
    # The run's codebook text said 60+; anyone excluded as under 60 is also under 65, so every age reason is reported
    # under the professor's rule.
    names = {r.id: (i4 if r.id == "I4" else f"{r.id} {r.name}") for r in codebook.rules}
    rows = collect(run_dir, codebook)
    title, abstract, fulltext = ({r["pmid"]: r for r in rows[k]} for k in ("title", "abstract", "fulltext"))
    sl = _read(run_dir / "second_look.csv")
    age_ft = _read(run_dir / "age_check.csv")
    age_ab = _read(run_dir / "age_check_abstract.csv")
    found = _read(free_found or run_dir / "free_fulltext_found.csv")

    for stage_rows in (title, abstract, fulltext):
        for r in stage_rows.values():
            if r["main reason"].startswith("I4 "):
                r["main reason"] = i4

    out = []
    for pmid, t in title.items():
        rec = {"pmid": pmid, "title": t["title"], "year": t["year"], "link": t["link"], "stage": "", "outcome": "",
               "reason": "", "decided by": "", "evidence": ""}
        if t["ai_outcome"] == "excluded":
            out.append(rec | {"stage": "title", "outcome": "excluded", "reason": t["main reason"],
                              "decided by": "panel: all 3 reviewers excluded", "evidence": t["quotes"]})
            continue
        a = abstract.get(pmid)
        if a is None or a["ai_outcome"] == "still open":
            out.append(rec | {"stage": "abstract", "outcome": "not finished", "reason": "abstract not screened yet"})
            continue
        if a["ai_outcome"] == "excluded":
            how = "panel: " + (a["how decided"] or "") + (" (2-1, flagged)" if a["escalated"] else "")
            out.append(rec | {"stage": "abstract", "outcome": "excluded", "reason": a["main reason"],
                              "decided by": how, "evidence": a["quotes"]})
            continue
        ab = age_ab.get(pmid)
        # An abstract age under 65 excludes at the abstract stage, unless the full text (the fuller source) confirms
        # 65+: then the abstract's number is about a subgroup (e.g. a younger comparison group), not the sample.
        if ab and ab.get("abstract decision", "").startswith("excluded") and age_ft.get(pmid, {}).get("I4 at 65") != "met":
            out.append(rec | {"stage": "abstract", "outcome": "excluded", "reason": i4,
                              "decided by": "age check on the abstract", "evidence": f"{ab['reason']}: {ab['quotes']}"})
            continue
        f = fulltext.get(pmid)
        if f is None or f["ai_outcome"] == "still open":
            out.append(rec | {"stage": "full text", "outcome": "not finished",
                              "reason": "full-text review not finished (a reviewer's host rejects the paper)"})
            continue
        if f["ai_outcome"] == "not retrieved":
            hand = (found.get(pmid) or {}).get("outcome", "").startswith("free")
            out.append(rec | {"stage": "full text", "outcome": HAND if hand else PAID,
                              "reason": "free copy exists; the site blocks scripts" if hand else "no free full text",
                              "evidence": (found.get(pmid) or {}).get("link", "")})
            continue
        if f["ai_outcome"] == "excluded":
            out.append(rec | {"stage": "full text", "outcome": "excluded", "reason": f["main reason"],
                              "decided by": "panel: " + f["how decided"], "evidence": f["quotes"]})
            continue
        age = age_ft.get(pmid, {})
        after = age.get("after", "")
        s = sl.get(pmid, {})
        if after.startswith("included"):
            how = "second look" if "second look" in after else "panel: " + f["how decided"]
            out.append(rec | {"stage": "full text", "outcome": "included", "reason": "all rules met",
                              "decided by": f"{how}; age 65+ confirmed ({age.get('reason', '')})",
                              "evidence": age.get("quotes", "")})
        elif after == "excluded (I4 under 65)":
            out.append(rec | {"stage": "full text", "outcome": "excluded", "reason": i4,
                              "decided by": "age check on the full text", "evidence": f"{age['reason']}: {age['quotes']}"})
        elif after == "excluded (second look)":
            out.append(rec | {"stage": "full text", "outcome": "excluded", "reason": _failing_rule(s.get("rule marks", ""), names),
                              "decided by": "second look", "evidence": s.get("quotes", "")})
        else:
            if after.startswith("for a person (age unclear"):
                why = "age 65+ not confirmable from the paper (it was included under 60+)"
            else:
                other = [r for r in _open_rules(s.get("rule marks", "")) if r != "I4"]
                why = (f"rules unclear in the paper: {', '.join(names.get(r, r) for r in other)}" if other
                       else "the paper does not report the participants' age")
            out.append(rec | {"stage": "full text", "outcome": "for the professor", "reason": why,
                              "decided by": "panel, second look and age check could not settle it",
                              "evidence": s.get("reasoning", "")})
    return out


def summarize(final: list[dict]) -> dict:
    by = Counter((r["stage"], r["outcome"]) for r in final)
    reasons = {st: dict(Counter(r["reason"] for r in final if r["stage"] == st and r["outcome"] == "excluded").most_common())
               for st in ("title", "abstract", "full text")}
    t_ex, a_ex = by[("title", "excluded")], by[("abstract", "excluded")]
    ft = {k: by[("full text", k)] for k in ("included", "excluded", "for the professor", "not finished", PAID, HAND)}
    free_read = ft["included"] + ft["excluded"] + ft["for the professor"] + ft["not finished"]
    return {
        "records": len(final),
        "title": {"screened": len(final), "excluded": t_ex, "passed": len(final) - t_ex},
        "abstract": {"screened": len(final) - t_ex, "excluded": a_ex, "passed": len(final) - t_ex - a_ex,
                     "excluded_by_age_check": sum(1 for r in final if r["decided by"] == "age check on the abstract")},
        "fulltext_free": {"assessed": free_read, **{k: ft[k] for k in ("included", "excluded", "for the professor",
                                                                          "not finished")}},
        "side_note": {"free_copy_to_save_by_hand": ft[HAND], "paid_access_not_assessed": ft[PAID]},
        "exclusion_reasons": reasons,
        "for_the_professor_reasons": dict(Counter(r["reason"] for r in final if r["outcome"] == "for the professor").most_common()),
    }


def stage_table(s: dict) -> str:
    f = s["fulltext_free"]
    lines = [
        "| Stage | In | Excluded | Passed / included | For the professor |",
        "|---|---|---|---|---|",
        f"| 1. Titles | {s['title']['screened']:,} | {s['title']['excluded']:,} | {s['title']['passed']:,} | — |",
        f"| 2. Abstracts | {s['abstract']['screened']:,} | {s['abstract']['excluded']:,} | {s['abstract']['passed']:,} | — |",
        f"| 3. Full text (free) | {f['assessed']:,} | {f['excluded']:,} | **{f['included']:,}** | {f['for the professor']:,}"
        + (f" (+{f['not finished']} not finished)" if f["not finished"] else "") + " |",
        "",
        f"Side note (not counted): {s['side_note']['free_copy_to_save_by_hand']:,} papers have a free copy still to be "
        f"saved by hand; {s['side_note']['paid_access_not_assessed']:,} are paid access only and were not assessed.",
    ]
    return "\n".join(lines)


def figures(final: list[dict], s: dict, out: Path) -> None:
    from adrd_review.eval import report_figures as rf

    plt = rf._plt()
    f = s["fulltext_free"]
    groups = {"Titles": [("Passed", s["title"]["passed"], rf.BLUE), ("Excluded", s["title"]["excluded"], rf.ORANGE)],
              "Abstracts": [("Passed", s["abstract"]["passed"], rf.BLUE), ("Excluded", s["abstract"]["excluded"], rf.ORANGE)],
              "Full text (free)": [("Included", f["included"], rf.BLUE), ("Excluded", f["excluded"], rf.ORANGE),
                                   ("For the professor", f["for the professor"], rf.AQUA),
                                   ("Not finished", f["not finished"], rf.GREY)]}
    fig, ax = plt.subplots(figsize=(9, 3.2))
    names = list(groups)[::-1]
    widest = max(sum(n for _, n, _ in groups[g]) for g in names)
    for y, g in enumerate(names):
        left = 0
        for label, n, color in groups[g]:
            if not n:
                continue
            ax.barh(y, n, left=left, color=color, height=0.55, edgecolor=rf.SURFACE, linewidth=2)
            if n >= widest * 0.04:
                ax.text(left + n / 2, y, f"{n:,}", ha="center", va="center", fontsize=9,
                        color="white" if color in (rf.BLUE, rf.ORANGE) else rf.INK)
            left += n
        small = [f"{n:,} {label.lower()}" for label, n, _ in groups[g] if n and n < widest * 0.04]
        ax.text(left + widest * 0.01, y, f"{left:,} papers" + (f"  ({', '.join(small)})" if small else ""),
                va="center", fontsize=9, color=rf.INK2)
    ax.set_yticks(range(len(names)), names)
    ax.set_xlim(0, widest * 1.3)
    ax.xaxis.grid(True, color=rf.GRID, linewidth=0.8)
    ax.set_axisbelow(True)
    ax.tick_params(axis="y", length=0)
    ax.legend([plt.Rectangle((0, 0), 1, 1, color=c) for c in (rf.BLUE, rf.ORANGE, rf.AQUA)],
              ["Passed / included", "Excluded", "For the professor"], ncol=3, frameon=False, loc="upper left",
              bbox_to_anchor=(0, -0.12), fontsize=9)
    ax.set_title("Final result: papers at each stage (older adults 65+, free full text only)")
    rf._save(fig, out, "final_stages", [{"stage": g, "outcome": l, "papers": n} for g in groups for l, n, _ in groups[g]])

    reasons = s["exclusion_reasons"]
    rules = sorted({r for d in reasons.values() for r in d}, key=lambda r: (r[0] != "I", r))
    fig, axes = plt.subplots(1, 3, figsize=(12, 0.45 * len(rules) + 1.4), sharey=True)
    top = max((v for d in reasons.values() for v in d.values()), default=1)
    data = []
    for ax, (st, label) in zip(axes, (("title", "Titles"), ("abstract", "Abstracts"), ("full text", "Full text (free)"))):
        vals = [reasons[st].get(r, 0) for r in rules]
        bars = ax.barh(range(len(rules)), vals, color=rf.ORANGE, height=0.6)
        for b, v in zip(bars, vals):
            if v:
                ax.text(b.get_width() + top * 0.02, b.get_y() + b.get_height() / 2, str(v), va="center", fontsize=9)
        ax.set_xlim(0, top * 1.25)
        ax.set_title(f"{label}: {sum(vals):,} excluded")
        ax.xaxis.grid(True, color=rf.GRID, linewidth=0.8)
        ax.set_axisbelow(True)
        ax.tick_params(axis="y", length=0)
        data += [{"stage": label, "rule": r, "excluded": v} for r, v in zip(rules, vals)]
    axes[0].set_yticks(range(len(rules)), [r if len(r) < 36 else r[:34] + "…" for r in rules])
    axes[0].invert_yaxis()
    fig.suptitle("Final: main reason for each exclusion", x=0.01, ha="left", fontweight="bold", fontsize=11, y=1.02)
    rf._save(fig, out, "final_reasons", data)


CHECK_GROUPS = [("included", lambda r: r["outcome"] == "included"),
                ("excluded at title", lambda r: r["stage"] == "title" and r["outcome"] == "excluded"),
                ("excluded at abstract", lambda r: r["stage"] == "abstract" and r["outcome"] == "excluded"),
                ("excluded at full text", lambda r: r["stage"] == "full text" and r["outcome"] == "excluded")]


def professor_check(final: list[dict], run_dir: Path, out: Path, n: int = 10, seed: int = 20261021) -> list[dict]:
    """A blind random check: n papers from each group (included; excluded at title / abstract / full text), shuffled.
    The professor's file shows only the paper; the AI's decision is in a separate answer key (and hidden columns)."""
    import random

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    from adrd_review.eval.stages import load_states

    abstracts = {s["pmid"]: s["article"].get("abstract") or "" for s in load_states(run_dir)}
    rng = random.Random(seed)
    picked = []
    for name, keep in CHECK_GROUPS:
        group = sorted((r for r in final if keep(r)), key=lambda r: r["pmid"])
        picked += [r | {"group": name} for r in rng.sample(group, min(n, len(group)))]
    rng.shuffle(picked)
    rows = [{"#": i, "PMID": r["pmid"], "PubMed link": r["link"], "title": r["title"], "year": r["year"],
             "abstract": abstracts.get(r["pmid"], ""), "your decision (include / exclude)": "",
             "your reason (rule, e.g. I1, I4, E1)": "", "notes": "",
             "AI outcome": f"{r['outcome']} ({r['stage']})" if r["outcome"] == "excluded" else r["outcome"],
             "AI reason": r["reason"], "AI decided by": r["decided by"], "AI evidence": r["evidence"]}
            for i, r in enumerate(picked, 1)]
    blind = [c for c in rows[0] if not c.startswith("AI ")]
    with (out / "professor_check.csv").open("w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=blind, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)
    with (out / "professor_check_answer_key.csv").open("w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=["#", "PMID", "title", "AI outcome", "AI reason", "AI decided by", "AI evidence"],
                           extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)

    wb = Workbook()
    ws = wb.active
    ws.title = "check (blind)"
    cols = list(rows[0])
    ws.append(cols)
    for r in rows:
        ws.append([r[c] for c in cols])
    widths = {"#": 5, "PMID": 11, "PubMed link": 18, "title": 45, "year": 6, "abstract": 80,
              "your decision (include / exclude)": 16, "your reason (rule, e.g. I1, I4, E1)": 16, "notes": 30}
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="FBEFD5" if c.startswith("your") or c == "notes" else "DCEFED")
        ws.column_dimensions[cell.column_letter].width = widths.get(c, 40)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        row[2].hyperlink = row[2].value
    first_ai = cols.index("AI outcome") + 1
    ws.column_dimensions.group(ws.cell(row=1, column=first_ai).column_letter,
                               ws.cell(row=1, column=len(cols)).column_letter, hidden=True)
    dv = DataValidation(type="list", formula1='"include,exclude,unsure"', allow_blank=True)
    ws.add_data_validation(dv)
    col = ws.cell(row=1, column=cols.index("your decision (include / exclude)") + 1).column_letter
    dv.add(f"{col}2:{col}{len(rows) + 1}")
    ws.freeze_panes = "D2"
    how = wb.create_sheet("how to check")
    for line in [
        f"{len(rows)} papers drawn at random (seed {seed}): {n} included, {n} excluded at title, {n} at abstract, "
        f"{n} at full text, in random order. Free full text only.",
        "For each paper, decide from the abstract and, if needed, the full paper (PubMed link) whether it belongs in "
        "the review under your criteria: include or exclude, and the main rule.",
        "The AI's decision is in the hidden columns on the right (click the + above them) and in "
        "professor_check_answer_key.csv: look only after deciding.",
        "Criteria: I1 dementia/Alzheimer's (MCI only together with it); I2 AI/ML (logistic regression counts); "
        "I3 detection/diagnosis/monitoring/prediction; I4 older adults 65+; E1 review; E2 abstract only; "
        "E3 no original results; E4 not human.",
    ]:
        how.append([line])
    how.column_dimensions["A"].width = 120
    for c in how["A"]:
        c.alignment = Alignment(wrap_text=True)
    wb.save(out / "professor_check.xlsx")
    return rows


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(prog="python -m adrd_review.final_results")
    ap.add_argument("run_dir")
    ap.add_argument("--free-found", default=None, help="free_fulltext_found.csv if it is in another run folder")
    args = ap.parse_args(argv)
    run_dir = Path(args.run_dir)
    final = build(run_dir, Path(args.free_found) if args.free_found else None)
    out = run_dir / "final"
    out.mkdir(exist_ok=True)
    cols = ["pmid", "title", "year", "link", "stage", "outcome", "reason", "decided by", "evidence"]
    for name, rows in (("decisions.csv", final),
                       ("for_professor.csv", [r for r in final if r["outcome"] == "for the professor"]),
                       ("not_assessed.csv", [r for r in final if r["outcome"].startswith("not assessed")])):
        with (out / name).open("w", newline="") as fh:
            w = csv.DictWriter(fh, fieldnames=cols)
            w.writeheader()
            w.writerows(rows)
    s = summarize(final)
    (out / "summary.json").write_text(json.dumps(s, indent=2) + "\n")
    table = stage_table(s)
    (out / "stage_table.md").write_text(table + "\n")
    figures(final, s, out)
    professor_check(final, run_dir, out)
    print(table)
    print(f"\nSaved in {out}/")
    return 0


if __name__ == "__main__":
    sys.exit(main())
