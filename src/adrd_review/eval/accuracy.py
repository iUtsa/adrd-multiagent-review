"""Accuracy and reliability for the professor's search.

There is no human gold standard for the 1,855 search results, so accuracy proper (agreement with a person) comes
from the professor's blind random check; until she fills it in, this module reports what can be measured honestly:

    reliability     agreement among the three AI reviewers per stage: Fleiss' kappa, share of papers where all
                    three agree, pairwise Cohen's kappa and % agreement (independent votes only)
    grounding       share of the reviewers' quotes found word for word in the text (unverifiable quotes are
                    discarded and the rule counts as unclear, so no decision rests on them)
    calibration     the one accuracy test against real human decisions so far (Cochrane dementia screening,
                    CLEF 2017, 40 papers): a different topic and codebook, reported as context only
    random check    when the professor's professor_check.xlsx (or .csv) is filled in: accuracy, sensitivity
                    (recall of includes), specificity, precision, NPV and Cohen's kappa, each with a 95% Wilson
                    interval, overall and per group

    python -m adrd_review.eval.accuracy data/runs/PROF_search_free/<timestamp>                   # reliability
    python -m adrd_review.eval.accuracy data/runs/PROF_search_free/<timestamp> --check <filled xlsx|csv>
Writes <run>/final/accuracy.json (+ accuracy_check.json when a check is scored).
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import sys
from collections import Counter
from itertools import combinations
from pathlib import Path

from adrd_review.eval.metrics import cohen_kappa, fleiss_kappa
from adrd_review.eval.stages import load_states

CALIBRATION = {
    "dataset": "Cochrane dementia screening reviews (CLEF 2017 TAR), calibration split, 40 papers with human decisions",
    "condition": "E2: 3-reviewer panel (Groq free models), codebook v0.2.0, abstract stage",
    "recall": 0.769, "precision": 0.769, "specificity": 0.889, "cohen_kappa": 0.658,
    "keyword_baseline": {"recall": 0.769, "precision": 0.435, "specificity": 0.519, "cohen_kappa": 0.240},
    "note": "A different topic and codebook from the professor's review: context, not an estimate for this search. "
            "The 3 missed papers were split votes, and Cochrane excluded all 3 after full text.",
}


def wilson(k: int, n: int, z: float = 1.96) -> list[float] | None:
    if n == 0:
        return None
    p = k / n
    d = 1 + z * z / n
    c = (p + z * z / (2 * n)) / d
    h = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / d
    return [round(max(0.0, c - h), 3), round(min(1.0, c + h), 3)]


def reliability(run_dir: Path) -> dict:
    states = load_states(run_dir)
    out = {}
    for field, name in (("title_votes", "titles"), ("screening_votes", "abstracts"), ("eligibility_votes", "full text")):
        rows = []
        for s in states:
            by = {v["reviewer_id"]: v["labels"][0] for v in s.get(field) or [] if int(v.get("round") or 0) == 0}
            if {"A", "B", "C"} <= by.keys():
                rows.append([by["A"], by["B"], by["C"]])
        if not rows:
            continue
        pairs = {}
        for i, j in combinations(range(3), 2):
            a, b = [r[i] for r in rows], [r[j] for r in rows]
            pairs["ABC"[i] + "-" + "ABC"[j]] = {"cohen_kappa": round(cohen_kappa(a, b), 3),
                                               "agreement": round(sum(x == y for x, y in zip(a, b)) / len(rows), 3)}
        out[name] = {"papers": len(rows), "fleiss_kappa": round(fleiss_kappa(rows), 3),
                     "all_three_agree": round(sum(len(set(r)) == 1 for r in rows) / len(rows), 3), "pairs": pairs,
                     "votes": dict(Counter(x for r in rows for x in r))}
    quotes = bad = 0
    for s in states:
        for field in ("title_votes", "screening_votes", "eligibility_votes"):
            for v in s.get(field) or []:
                quotes += len(v.get("evidence") or [])
                bad += sum("quote not found" in w for w in v.get("warnings") or [])
    out["grounding"] = {"quotes_given": quotes + bad, "verified": quotes, "not_found_and_discarded": bad,
                        "verified_share": round(quotes / (quotes + bad), 3) if quotes + bad else None}
    return out


def kappa_label(k: float) -> str:  # Landis & Koch
    return ("poor" if k < 0 else "slight" if k <= 0.2 else "fair" if k <= 0.4 else "moderate" if k <= 0.6
            else "substantial" if k <= 0.8 else "almost perfect")


def score_check(check: Path, key: Path) -> dict:
    """Score the professor's filled random check against the AI's answers (include vs not include)."""
    if check.suffix == ".xlsx":
        from openpyxl import load_workbook

        ws = load_workbook(check)["check (blind)"]
        head = [c.value for c in ws[1]]
        rows = [dict(zip(head, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
    else:
        rows = list(csv.DictReader(check.open()))
    ai = {str(r["#"]): r for r in csv.DictReader(key.open())}
    pairs = []
    for r in rows:
        d = str(r.get("your decision (include / exclude)") or "").strip().lower()
        if d in ("include", "exclude"):
            a = ai[str(r["#"])]
            pairs.append((d == "include", a["AI outcome"] == "included", a["AI outcome"]))
    n = len(pairs)
    if not n:
        return {"scored": 0, "note": "no include/exclude decisions filled in yet"}
    tp = sum(h and m for h, m, _ in pairs)
    tn = sum(not h and not m for h, m, _ in pairs)
    fp = sum(not h and m for h, m, _ in pairs)
    fn = sum(h and not m for h, m, _ in pairs)

    def stat(k, d):
        return {"value": round(k / d, 3) if d else None, "ci95": wilson(k, d), "n": d}

    groups = Counter(g for _, _, g in pairs)
    return {
        "scored": n, "unsure_or_blank": len(rows) - n,
        "accuracy": stat(tp + tn, n), "sensitivity_recall": stat(tp, tp + fn), "specificity": stat(tn, tn + fp),
        "precision": stat(tp, tp + fp), "npv": stat(tn, tn + fn),
        "cohen_kappa": round(cohen_kappa([h for h, _, _ in pairs], [m for _, m, _ in pairs]), 3),
        "confusion": {"both include": tp, "both exclude": tn, "AI include, professor exclude": fp,
                      "AI exclude, professor include (missed)": fn},
        "agreement_by_AI_group": {g: round(sum(h == m for h, m, x in pairs if x == g) / c, 3) for g, c in groups.items()},
        "note": "The sample has 10 papers per group, so it is not proportional to the screened set: per-group agreement "
                "is the direct reading; the overall figures weight the groups equally.",
    }


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(prog="python -m adrd_review.eval.accuracy")
    ap.add_argument("run_dir")
    ap.add_argument("--check", default=None, help="the professor's filled professor_check.xlsx or .csv")
    args = ap.parse_args(argv)
    run_dir = Path(args.run_dir)
    out = run_dir / "final"
    if args.check:
        res = score_check(Path(args.check), out / "professor_check_answer_key.csv")
        (out / "accuracy_check.json").write_text(json.dumps(res, indent=2) + "\n")
        print(json.dumps(res, indent=2))
        return 0
    res = {"reliability": reliability(run_dir), "calibration_human_gold": CALIBRATION}
    (out / "accuracy.json").write_text(json.dumps(res, indent=2) + "\n")
    print(json.dumps(res, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
