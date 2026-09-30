"""Error analysis (task 10.2): export every disagreement with the gold standard to a
spreadsheet for a person to categorise, then count the categories.

    python -m adrd_review.eval.errors export data/runs/E4_gemini/<timestamp>
    python -m adrd_review.eval.errors summarize data/runs/E4_gemini/<timestamp>/error_analysis.xlsx

Categories follow the causes of LLM-human disagreement found in prior work
(Understanding LLMs in title-abstract screening, 2026) plus two practical ones.
"""

from __future__ import annotations

import argparse
import csv
import json
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from adrd_review.codebook import load_codebook
from adrd_review.gold import load_gold
from adrd_review.tools.pubmed import load_cached

CATEGORIES = [
    "unclear rule (boundary ambiguity)",
    "keyword trap (matched a word, missed the meaning)",
    "wrong topic inference",
    "missing information in abstract",
    "gold label questionable",
    "other",
]

COLUMNS = [
    ("pmid", 11), ("title", 50), ("year", 6), ("field", 11), ("gold", 16), ("agent", 16),
    ("how decided", 13), ("escalated", 16), ("reviewer votes", 30), ("reasoning", 60), ("quotes", 60),
    ("adjudicator", 45), ("abstract", 70), ("hint", 22), ("error_category", 34), ("notes", 40),
]


def _labels(x) -> str:
    return ", ".join(x) if x else "(none)"


def collect(run_dir: Path) -> list[dict]:
    meta = json.loads((run_dir / "run_meta.json").read_text())
    cfg = meta["config"]
    codebook = load_codebook(cfg["codebook"])
    gold = {g.pmid: g for g in load_gold(cfg["gold"], codebook)}
    rows = []
    for line in (run_dir / "decisions.jsonl").read_text().splitlines():
        d = json.loads(line)
        g = gold.get(d["pmid"])
        if g is None:
            continue
        art = load_cached(d["pmid"], cfg.get("corpus_dir", "data/corpus"))
        agent_scr = d["screening_decision"]["labels"]
        checks = [("screening", [g.screening], agent_scr, d["screening_decision"]["method"])]
        if g.screening == "include":
            for dim in codebook.dimensions:
                truth = g.labels.get(dim, [])
                if not truth:
                    continue  # placeholder gold has positives-only labels; nothing to compare
                dec = (d.get("annotation_decisions") or {}).get(dim) if agent_scr == ["include"] else None
                checks.append((dim, truth, dec["labels"] if dec else [], dec["method"] if dec else "not annotated"))
        for field, truth, agent, method in checks:
            if sorted(truth) == sorted(agent):
                continue
            votes = [v for v in d.get("screening_votes" if field == "screening" else "annotation_votes", [])
                     if v["field"] == field]
            latest = {}
            for v in sorted(votes, key=lambda v: v["round"]):
                latest[v["reviewer_id"]] = v
            adj = [a for a in d.get("adjudications", []) if a["field"] == field]
            esc = sorted({e["reason"] for e in d.get("escalations", []) if e["field"] == field})
            no_info = (art is not None and not art.has_abstract) or (
                field == "screening" and d.get("triage") and not d["triage"]["enough_information"]
            )
            rows.append({
                "pmid": d["pmid"],
                "title": art.title if art else "",
                "year": art.year if art else "",
                "field": field,
                "gold": _labels(truth),
                "agent": _labels(agent),
                "how decided": method,
                "escalated": ", ".join(esc),
                "reviewer votes": "; ".join(
                    f"{r}: {_labels(v['labels'])}" + (f" [{', '.join(v['rule_ids'])}]" if v["rule_ids"] else "")
                    for r, v in sorted(latest.items())
                ),
                "reasoning": "\n".join(f"{r}: {v['rationale']}" for r, v in sorted(latest.items()) if v["rationale"]),
                "quotes": "\n".join(f'{r}: "{e["quote"]}"' for r, v in sorted(latest.items()) for e in v["evidence"]),
                "adjudicator": "\n".join(f"{a['rationale']} (confidence {a['confidence']})" for a in adj),
                "abstract": (art.abstract if art and art.has_abstract else "(no abstract)")[:3000],
                "hint": "missing information?" if no_info else "",
                "error_category": "",
                "notes": "",
            })
    return rows


def export(run_dir: Path, out: Path | None = None) -> Path:
    rows = collect(run_dir)
    out = out or run_dir / "error_analysis.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "disagreements"
    ws.append([c for c, _ in COLUMNS])
    for r in rows:
        ws.append([r[c] for c, _ in COLUMNS])
    head_fill = PatternFill("solid", fgColor="DCEFED")
    for i, (_, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font, cell.fill = Font(bold=True), head_fill
        ws.column_dimensions[cell.column_letter].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    cat_col = ws.cell(row=1, column=[c for c, _ in COLUMNS].index("error_category") + 1).column_letter
    dv = DataValidation(type="list", formula1='"' + ",".join(CATEGORIES) + '"', allow_blank=True)
    dv.error, dv.prompt = "Pick a category from the list.", "Why did the agents disagree with the gold standard?"
    ws.add_data_validation(dv)
    dv.add(f"{cat_col}2:{cat_col}{max(2, len(rows) + 1)}")

    guide = wb.create_sheet("how to code")
    guide.append(["Category", "Use when"])
    explanations = [
        "The rule's wording allows both answers; a clearer definition or example would settle it.",
        "The agents reacted to a word (e.g. an abbreviation such as MOCA) without its meaning.",
        "The agents misread what the study is about (population, test, or aim).",
        "The abstract does not contain the information needed; full text would be required.",
        "On reflection the agents' answer looks right and the gold label looks wrong.",
        "None of the above; explain in notes.",
    ]
    for c, e in zip(CATEGORIES, explanations):
        guide.append([c, e])
    guide.column_dimensions["A"].width, guide.column_dimensions["B"].width = 48, 90
    wb.save(out)
    return out


def summarize(sheet: Path) -> dict:
    ws = load_workbook(sheet)["disagreements"]
    header = [c.value for c in ws[1]]
    fi, ci = header.index("field"), header.index("error_category")
    counts: Counter = Counter()
    uncoded = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        cat = (row[ci] or "").strip()
        if not cat:
            uncoded += 1
        else:
            counts[(row[fi], cat)] += 1
    by_cat = Counter()
    for (_, cat), n in counts.items():
        by_cat[cat] += n
    return {"by_category": dict(by_cat.most_common()), "by_field_and_category": {f"{f} | {c}": n for (f, c), n in sorted(counts.items())},
            "uncoded": uncoded}


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(prog="python -m adrd_review.eval.errors")
    sub = ap.add_subparsers(dest="cmd", required=True)
    e = sub.add_parser("export", help="Write error_analysis.xlsx for a run")
    e.add_argument("run_dir")
    e.add_argument("--out")
    s = sub.add_parser("summarize", help="Count categories in a coded sheet")
    s.add_argument("sheet")
    args = ap.parse_args(argv)
    if args.cmd == "export":
        run_dir = Path(args.run_dir)
        out = export(run_dir, Path(args.out) if args.out else None)
        n = len(collect(run_dir))
        print(f"{n} disagreements -> {out}\nFill in the 'error_category' column (drop-down), then run:\n"
              f"  python -m adrd_review.eval.errors summarize {out}")
    else:
        result = summarize(Path(args.sheet))
        print(json.dumps(result, indent=2))
        out = Path(args.sheet).with_name("error_summary.csv")
        with open(out, "w", newline="") as fh:
            w = csv.writer(fh)
            w.writerow(["field", "category", "count"])
            for key, n in result["by_field_and_category"].items():
                w.writerow([*key.split(" | ", 1), n])
        print(f"Saved: {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
