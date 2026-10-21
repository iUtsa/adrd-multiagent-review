"""Every paper's decision, one file per stage, for a person to check (e.g. the professor spot-checking).

    1_title.csv      every title screened: excluded or passed, each reviewer's rule marks and quotes
    2_abstract.csv   every abstract screened (title and abstract included)
    3_fulltext.csv   every paper that reached full text: not retrieved, excluded, passed or waiting for a person
    4_included.csv   the papers included at the end
    stages.xlsx      the same four sheets plus 'random sample' (a fixed-seed draw of excluded AND passed papers
                     from each stage) and 'how to check'

In the workbook the AI's decision columns are hidden, so a reviewer can decide from the paper first
(human_decision) and then unhide them to compare (blind check).

Stages are assigned as in the PRISMA flow (eval/flow.py). Papers still open (e.g. waiting for a daily limit)
are included from in_progress.jsonl: listed at each stage they finished, and as 'still open' at the stage they
wait for. The flow counts them only once they are finished, so a stage file can list more papers than the flow.

    python -m adrd_review.eval.stages data/runs/<name>/<timestamp>
    python -m adrd_review.eval.stages data/runs/<name>/<timestamp> --sample 10 --seed 7
"""

from __future__ import annotations

import argparse
import csv
import json
import random
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from adrd_review.codebook import Codebook, load_codebook

DEFAULT_SAMPLE = 5  # per outcome per stage: 5 excluded + 5 passed at each stage
DEFAULT_SEED = 20261017
NOT_RETRIEVED = "not_retrieved"

STAGES = [("title", "1_title", "1 titles"), ("abstract", "2_abstract", "2 abstracts"),
          ("fulltext", "3_fulltext", "3 full text"), ("included", "4_included", "4 included")]

# The paper and the checker's columns come first; the AI's columns follow (hidden in the workbook).
PAPER_COLS = [("stage", 10), ("pmid", 11), ("link", 18), ("title", 50), ("year", 6), ("journal", 24),
              ("abstract", 70), ("full text source", 14)]
CHECK_COLS = [("human_decision", 14), ("human_rule", 10), ("human_notes", 40)]
AI_COLS = [("ai_outcome", 16), ("how decided", 14), ("main reason", 34), ("escalated", 16), ("reviewer votes", 30),
           ("rule marks", 60), ("quotes", 60), ("reasoning", 60)]
COLUMNS = PAPER_COLS + CHECK_COLS + AI_COLS
DECISIONS = ["include", "exclude", "unsure"]

MARK = {"met": "met", "not_met": "NOT MET", "unclear": "unclear", "applies": "APPLIES", "does_not_apply": "no"}


def load_states(run_dir: Path) -> list[dict]:
    """Finished articles (decisions.jsonl) and what the still-open ones have finished so far (in_progress.jsonl)."""
    states = []
    for name in ("decisions.jsonl", "in_progress.jsonl"):
        if (run_dir / name).exists():
            states += [json.loads(line) for line in (run_dir / name).read_text().splitlines() if line.strip()]
    return states


def _labels(d: dict | None) -> list[str]:
    return (d or {}).get("labels") or []


def split_by_stage(states: list[dict]) -> dict[str, list[tuple[dict, str]]]:
    """(state, outcome) per stage, assigned as in eval.flow.flow_counts. An article still open (its state has
    'open') is listed at every stage it finished, and as 'still open' at the stage it is waiting for."""
    out: dict[str, list[tuple[dict, str]]] = {k: [] for k, _, _ in STAGES}
    for s in states:
        is_open = bool(s.get("open"))
        if s.get("title_decision"):
            if _labels(s["title_decision"]) == ["exclude"]:
                out["title"].append((s, "excluded"))
                continue
            out["title"].append((s, "passed"))

        if not s.get("screening_decision"):
            if is_open and s.get("title_decision"):
                out["abstract"].append((s, "still open"))
            continue
        if _labels(s["screening_decision"]) != ["include"]:
            out["abstract"].append((s, "excluded"))
            continue
        out["abstract"].append((s, "passed"))

        d = s.get("eligibility_decision") or {}
        reached = (s.get("fulltext_info") or {}).get("stage") == 2
        if is_open and not d:
            out["fulltext"].append((s, "still open"))
            continue
        if reached:
            if d.get("method") == NOT_RETRIEVED:
                out["fulltext"].append((s, "not retrieved"))
                continue
            if _labels(d) == ["exclude"]:
                out["fulltext"].append((s, "excluded"))
                continue
            out["fulltext"].append((s, "waiting for a person" if _labels(d) == ["unclear"] else "passed"))
        out["included"].append((s, "waiting for a person" if _labels(d) == ["unclear"] else "included"))
    return out


def _latest(votes: list[dict]) -> dict[str, dict]:
    latest: dict[str, dict] = {}
    for v in sorted(votes, key=lambda v: int(v.get("round") or 0)):
        latest[v["reviewer_id"]] = v
    return latest


def _main_reason(votes: dict[str, dict], decision: dict | None, names: dict[str, str]) -> str:
    """The rule cited by most of the reviewers who excluded (codebook order breaks ties), as in eval.flow."""
    if decision and decision.get("method") == "human":
        return "decided by a person"
    counts = Counter(r for v in votes.values() if v.get("labels") == ["exclude"] for r in v.get("rule_ids") or [])
    if not counts:
        return ""
    order = {r: i for i, r in enumerate(names)}
    best = min(counts, key=lambda r: (-counts[r], order.get(r, 99), r))
    return names.get(best, best)


def _row(s: dict, stage: str, outcome: str, names: dict[str, str]) -> dict:
    a = s["article"]
    field = {"title": "title", "abstract": "screening", "fulltext": "eligibility"}.get(stage)
    if stage == "included":  # the last stage that judged it
        field = "eligibility" if s.get("eligibility_votes") else "screening"
    votes = _latest(s.get(f"{field}_votes") or [])
    decision = s.get(f"{field}_decision")
    info = s.get("fulltext_info") or {}
    esc = sorted({e.get("reason", "") for e in s.get("escalations") or [] if e.get("field") == field})
    return {
        "stage": stage,
        "pmid": s["pmid"],
        "link": f"https://pubmed.ncbi.nlm.nih.gov/{s['pmid']}/",
        "title": a.get("title") or "",
        "year": a.get("year") or "",
        "journal": a.get("journal") or "",
        "abstract": "" if stage == "title" else (a.get("abstract") or "(no abstract)"),
        "full text source": (info.get("source") or "none") if info.get("stage") == 2 else "",
        "human_decision": "", "human_rule": "", "human_notes": "",
        "ai_outcome": outcome,
        "how decided": (decision or {}).get("method") or ("open: " + s["open"][:80] if outcome == "still open" else ""),
        "main reason": (_main_reason(votes, decision, names) if outcome == "excluded"
                        else "no free full text" if outcome == "not retrieved" else ""),
        "escalated": ", ".join(esc),
        "reviewer votes": "; ".join(
            f"{r}: {', '.join(v.get('labels') or []) or '(none)'}" + (f" [{', '.join(v['rule_ids'])}]" if v.get("rule_ids") else "")
            for r, v in sorted(votes.items())),
        "rule marks": "\n".join(
            f"{r}: " + " · ".join(f"{k} {MARK.get(m, m)}" for k, m in (v.get("rule_checks") or {}).items())
            for r, v in sorted(votes.items()) if v.get("rule_checks")),
        "quotes": "\n".join(f'{r}: "{e["quote"]}"' for r, v in sorted(votes.items()) for e in v.get("evidence") or []
                            if e.get("quote")),
        "reasoning": "\n".join(f"{r}: {v['rationale']}" for r, v in sorted(votes.items()) if v.get("rationale")),
    }


def collect(run_dir: Path, codebook: Codebook | None = None) -> dict[str, list[dict]]:
    if codebook is None and (run_dir / "run_meta.json").exists():
        path = json.loads((run_dir / "run_meta.json").read_text()).get("config", {}).get("codebook")
        codebook = load_codebook(path) if path and Path(path).exists() else None
    names = {r.id: f"{r.id} {r.name}" for r in codebook.rules} if codebook else {}
    by_stage = split_by_stage(load_states(run_dir))
    return {stage: [_row(s, stage, outcome, names) for s, outcome in pairs] for stage, pairs in by_stage.items()}


def sample(rows: dict[str, list[dict]], n: int = DEFAULT_SAMPLE, seed: int = DEFAULT_SEED) -> list[dict]:
    """Up to n decided papers per (stage, outcome), drawn with a fixed seed so the same sample can be drawn again."""
    rng = random.Random(seed)
    picked = []
    for stage, _, _ in STAGES:
        groups: dict[str, list[dict]] = {}
        for r in sorted(rows[stage], key=lambda r: r["pmid"]):
            # Nothing to check: not decided yet, or no free full text (paid-access papers are not assessed and are
            # left out of every accuracy measure; their number is only reported alongside).
            if r["ai_outcome"] not in ("still open", "not retrieved"):
                groups.setdefault(r["ai_outcome"], []).append(r)
        for outcome in sorted(groups):
            picked += rng.sample(groups[outcome], min(n, len(groups[outcome])))
    return picked


def _sheet(wb: Workbook, title: str, rows: list[dict]) -> None:
    ws = wb.create_sheet(title)
    ws.append([c for c, _ in COLUMNS])
    for r in rows:
        ws.append([r[c] for c, _ in COLUMNS])
    for i, (name, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="FBEFD5" if (name, width) in CHECK_COLS else "DCEFED")
        ws.column_dimensions[cell.column_letter].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "D2"
    first_ai = len(PAPER_COLS) + len(CHECK_COLS) + 1
    ws.column_dimensions.group(ws.cell(row=1, column=first_ai).column_letter,
                               ws.cell(row=1, column=len(COLUMNS)).column_letter, hidden=True)
    col = ws.cell(row=1, column=[c for c, _ in COLUMNS].index("human_decision") + 1).column_letter
    dv = DataValidation(type="list", formula1='"' + ",".join(DECISIONS) + '"', allow_blank=True)
    dv.prompt = "Human decision from the paper alone, before unhiding the AI columns."
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}{max(2, len(rows) + 1)}")


def export(run_dir: Path, out_dir: Path | None = None, n: int = DEFAULT_SAMPLE, seed: int = DEFAULT_SEED,
           codebook: Codebook | None = None) -> Path:
    rows = collect(run_dir, codebook)
    out_dir = out_dir or run_dir / "stages"
    out_dir.mkdir(parents=True, exist_ok=True)
    for stage, file, _ in STAGES:
        with (out_dir / f"{file}.csv").open("w", newline="") as fh:
            w = csv.DictWriter(fh, fieldnames=[c for c, _ in COLUMNS])
            w.writeheader()
            w.writerows(rows[stage])

    wb = Workbook()
    wb.remove(wb.active)
    picked = sample(rows, n, seed)
    _sheet(wb, "random sample", picked)
    for stage, _, title in STAGES:
        _sheet(wb, title, rows[stage])
    failures = run_dir / "failures.json"
    still_open = len(json.loads(failures.read_text())) if failures.exists() else 0
    in_files = sum(1 for s in load_states(run_dir) if s.get("open"))
    guide = wb.create_sheet("how to check")
    for line in [
        ["Run", str(run_dir)],
        ["Random sample", f"up to {n} papers per stage and outcome (excluded, passed, ...), seed {seed}: "
                          f"python -m adrd_review.eval.stages {run_dir} --sample {n} --seed {seed} draws the same papers"],
        ["Blind check", "Decide from the paper first (human_decision: include / exclude / unsure; human_rule: e.g. I1 or E1), "
                        "then click the + above the columns on the right to show the AI's decision and compare."],
        ["What each stage saw", "Titles: the title only. Abstracts: title and abstract. Full text: the paper "
                                "(PubMed Central or a PDF in data/pdfs)."],
        ["Free full text only", "Papers whose full text is not freely available are not assessed and are not part of "
                                "any accuracy measure; '3 full text' lists them as 'not retrieved' for the record only."],
        ["Rule marks", "Per reviewer: inclusion rules met / NOT MET / unclear; exclusion rules APPLIES / no. "
                       "A paper is excluded only when a rule is clearly NOT MET or APPLIES, with a quote."],
        ["Panel", "Titles: excluded only if every reviewer excludes. Abstracts and full text: majority, "
                  "split exclusions go to a person."],
        ["Still open", f"{still_open} paper(s) not finished yet (waiting for a daily limit, a key or a PDF); {in_files} "
                       "of them are listed at the stages they finished, and as 'still open' where they wait."],
    ]:
        guide.append(line)
    guide.column_dimensions["A"].width, guide.column_dimensions["B"].width = 22, 110
    for row in guide.iter_rows():
        row[0].font = Font(bold=True)
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(out_dir / "stages.xlsx")
    return out_dir


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(prog="python -m adrd_review.eval.stages")
    ap.add_argument("run_dir")
    ap.add_argument("--out", help="folder to write to (default: <run_dir>/stages)")
    ap.add_argument("--sample", type=int, default=DEFAULT_SAMPLE, help="papers per stage and outcome in the random sample")
    ap.add_argument("--seed", type=int, default=DEFAULT_SEED)
    args = ap.parse_args(argv)
    run_dir = Path(args.run_dir)
    out = export(run_dir, Path(args.out) if args.out else None, args.sample, args.seed)
    rows = collect(run_dir)
    for stage, file, _ in STAGES:
        outcomes = Counter(r["ai_outcome"] for r in rows[stage])
        print(f"  {file}.csv  {len(rows[stage]):5} papers   " + ", ".join(f"{k} {v}" for k, v in sorted(outcomes.items())))
    print(f"Saved in {out}/ (stages.xlsx: one sheet per stage plus a random sample, seed {args.seed})")
    return 0


if __name__ == "__main__":
    sys.exit(main())
