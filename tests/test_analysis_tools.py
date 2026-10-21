"""Aggregation, error sheet, figures, retrieval check, weights, cost and gold conversion."""

import json
import math
from pathlib import Path

import pandas as pd
import pytest
import yaml

from adrd_review.convert_gold import convert
from adrd_review.eval import aggregate, cost, errors, figures, retrieval, weights

ROOT = Path(__file__).resolve().parents[1]


# ---- helpers: a tiny finished run on disk -------------------------------------------


def make_run(root: Path, name: str, stamp: str, preds: dict[str, str], recall: float, repeat=0, part="calibration",
             gold="data/gold/gold.csv", decisions=None, usage=None):
    d = root / name / stamp
    d.mkdir(parents=True)
    pmids = sorted(preds)
    meta = {"name": name, "condition": name.split("_")[0], "part": part, "pmids": pmids, "limit": None, "repeat": repeat,
            "started_utc": stamp, "codebook_version": "t", "codebook_sha256": "c", "gold_sha256": "g",
            "prompts_sha256": "p", "config": {"name": name, "gold": gold, "codebook": "configs/codebook.yaml"}}
    (d / "run_meta.json").write_text(json.dumps(meta))
    report = {"screening": {"recall": recall, "precision": 0.5, "f1": 0.6, "specificity": 0.7, "cohen_kappa": 0.3,
                            "wss": 0.2, "wss_at_95": 0.1}, "annotation": {}, "escalations": {"rate": 0.25},
              "adjudications": 2, "usage": {"calls": 10, "per_article": {"input_tokens": 100, "output_tokens": 10},
                                            "total": {"cost_usd": 0}, "wall_seconds": 1}, "failures": 0}
    (d / "report.json").write_text(json.dumps(report))
    (d / "predictions.csv").write_text("pmid,screening,score\n" + "".join(f"{p},{s},0.5\n" for p, s in preds.items()))
    (d / "decisions.jsonl").write_text("\n".join(json.dumps(x) for x in (decisions or [])))
    (d / "usage.jsonl").write_text("\n".join(json.dumps(x) for x in (usage or [])))
    return d


# ---- aggregation --------------------------------------------------------------------


def test_aggregate_averages_repeats_and_measures_consistency(tmp_path):
    make_run(tmp_path, "E2_x", "20260101-000001-r1", {"1": "include", "2": "exclude"}, 0.8, repeat=1)
    make_run(tmp_path, "E2_x", "20260101-000002-r2", {"1": "include", "2": "include"}, 0.6, repeat=2)
    rows = aggregate.summarize(aggregate.latest_groups(aggregate.find_runs(tmp_path, "calibration")))
    (row,) = rows
    assert row["runs"] == 2 and row["recall_mean"] == pytest.approx(0.7)
    assert row["recall_sd"] == pytest.approx(0.1414, abs=1e-4)
    assert row["includes_not_stable"] == 1


def test_aggregate_never_mixes_runs_from_different_code(tmp_path):
    old = make_run(tmp_path, "E3_x", "20260101-000001", {"1": "include"}, 0.1)
    meta = json.loads((old / "run_meta.json").read_text()) | {"prompts_sha256": "OLD"}
    (old / "run_meta.json").write_text(json.dumps(meta))
    make_run(tmp_path, "E3_x", "20260102-000001", {"1": "include"}, 0.9)
    (row,) = aggregate.summarize(aggregate.latest_groups(aggregate.find_runs(tmp_path, "calibration")))
    assert row["runs"] == 1 and row["recall_mean"] == 0.9  # only the newer group


# ---- weights ---------------------------------------------------------------------------


def test_log_odds_weights():
    assert weights.log_odds(0.9) == pytest.approx(math.log(9))
    assert weights.log_odds(0.3) == weights.log_odds(0.55)  # clipped: never negative or zero
    assert weights.log_odds(0.999) == weights.log_odds(0.95)


def test_weights_refuse_test_runs(tmp_path):
    run = make_run(tmp_path, "E2_x", "20260101-000001", {"1": "include"}, 0.8, part="test")
    with pytest.raises(SystemExit, match="calibration runs only"):
        weights.reviewer_weights(run)


# ---- cost ------------------------------------------------------------------------------


def test_cost_scales_and_prices_by_role(tmp_path):
    usage = [{"role": "screening", "model": "google:a", "input_tokens": 1000, "output_tokens": 100}] * 4
    run = make_run(tmp_path, "E1_x", "20260101-000001", {"1": "include", "2": "exclude"}, 0.8, usage=usage)
    prices = {"google:a": {"input": 1.0, "output": 10.0}}
    r = cost.estimate(run, prices, articles=20, repeats=3)
    per_call = (1000 * 1.0 + 100 * 10.0) / 1e6
    assert r["total_usd"] == pytest.approx(round(4 * per_call * (20 * 3 / 2), 2))
    assert r["by_role"][0]["calls_per_article"] == 2.0


def test_cost_reports_missing_prices(tmp_path):
    run = make_run(tmp_path, "E1_x", "20260101-000001", {"1": "include"}, 0.8,
                   usage=[{"role": "triage", "model": "google:unknown", "input_tokens": 1, "output_tokens": 1}])
    with pytest.raises(SystemExit, match="google:unknown"):
        cost.estimate(run, {}, 10, 1)


# ---- retrieval ------------------------------------------------------------------------


def test_query_limits_are_read_from_the_query():
    q = yaml.safe_load((ROOT / "configs/search/prof_query.yaml").read_text())["corrected"]
    assert retrieval.query_limits(q) == {"years": (2015, 2025), "language": "eng"}


# ---- error sheet + figures on real fake runs (skipped if none exist) --------------------

RUNS = sorted((ROOT / "data/runs/E3_fake").glob("*-r1")) if (ROOT / "data/runs/E3_fake").exists() else []


@pytest.mark.skipif(not RUNS, reason="needs a fake E3 repeat run")
def test_error_sheet_round_trip(tmp_path):
    out = errors.export(RUNS[-1], tmp_path / "sheet.xlsx")
    from openpyxl import load_workbook

    wb = load_workbook(out)
    ws = wb["disagreements"]
    header = [c.value for c in ws[1]]
    ws.cell(row=2, column=header.index("error_category") + 1, value=errors.CATEGORIES[0])
    wb.save(out)
    s = errors.summarize(out)
    assert s["by_category"] == {errors.CATEGORIES[0]: 1} and s["uncoded"] == ws.max_row - 2


@pytest.mark.skipif(not RUNS, reason="needs a fake E3 repeat run")
def test_effort_curve_goes_from_agents_to_all_human():
    rows = figures.effort_curve(RUNS[-1])
    assert rows[0]["reviewed"] == 0 and rows[-1]["reviewed_share"] == 1.0
    assert rows[-1]["f1"] == 1.0 and rows[-1]["kappa"] == 1.0  # every answer from the gold standard


# ---- gold conversion ---------------------------------------------------------------------


def test_convert_maps_values_looks_up_pmids_and_reports_unmapped(tmp_path):
    src = tmp_path / "prof.xlsx"
    pd.DataFrame({
        "PMID": ["111", "", "333"],
        "DOI": ["", "10.1/abc", ""],
        "Title": ["A", "B", "C"],
        "Final decision": ["Included", "Excluded", "Include"],
        "Screening instrument": ["Mini-Mental State Examination; MoCA", "", "SLUMS"],
        "Reviewer 1": ["Yes", "No", "Yes"],
    }).to_excel(src, index=False)
    mapping = {
        "file": str(src),
        "id": {"pmid": "PMID", "doi": "DOI", "title": "Title"},
        "screening": {"column": "Final decision", "values": {"Included": "include", "Include": "include",
                                                            "Excluded": "exclude", "Yes": "include", "No": "exclude"}},
        "dimensions": {"instrument": {"column": "Screening instrument", "separator": ";",
                                      "values": {"Mini-Mental State Examination": "MMSE", "MoCA": "MoCA"}}},
        "reviewers": [{"id": "r1", "screening": "Reviewer 1"}],
    }
    std, problems, unmapped = convert(mapping, lookup=lambda doi=None, title=None: "222" if doi == "10.1/abc" else None)
    assert problems == []
    assert list(std["pmid"]) == ["111", "222", "333"]
    assert list(std["screening"]) == ["include", "exclude", "include"]
    assert std.loc[0, "instrument"] == "MMSE;MoCA" and std.loc[1, "instrument"] == ""
    assert std.loc[0, "screening__r1"] == "include"
    assert unmapped == {("instrument", "SLUMS"): 1}


def test_convert_reports_missing_columns(tmp_path):
    src = tmp_path / "prof.csv"
    src.write_text("id,decision\n1,Yes\n")
    mapping = {"file": str(src), "id": {"pmid": "PMID"}, "screening": {"column": "Final decision", "values": {}}}
    _, problems, _ = convert(mapping)
    assert any("'PMID' not in file" in p for p in problems)


def test_report_figures_write_png_and_csv_and_read_the_loop_log(tmp_path):
    from adrd_review.eval import report_figures

    rows = {"title": [{"ai_outcome": "passed"}] * 3 + [{"ai_outcome": "excluded"}],
            "abstract": [{"ai_outcome": "passed"}] * 2 + [{"ai_outcome": "excluded"}],
            "fulltext": [{"ai_outcome": "passed"}, {"ai_outcome": "waiting for a person"}, {"ai_outcome": "not retrieved"}],
            "included": []}
    png = report_figures.stage_outcomes(rows, tmp_path)
    assert png.exists() and png.stat().st_size > 1000
    csv_rows = (tmp_path / "stage_outcomes.csv").read_text()
    assert "Full text (free),Included,1" in csv_rows and "not retrieved" not in csv_rows.lower()  # paid access never charted
    log = tmp_path / "loop.log"
    log.write_text("=== try 1: 2026-10-11 08:42 ===\n=== try 1 done 10:46: 1597 article(s) still open; run folder x ===\n"
                   "=== try 2: 2026-10-11 23:00 ===\n=== try 2 done 01:10: 1500 article(s) still open; run folder x ===\n")
    assert report_figures.progress(log, tmp_path).exists()
    assert "2026-10-12T01:10,1500" in (tmp_path / "progress.csv").read_text()  # finished after midnight


def test_random_check_scoring_counts_misses_and_gives_intervals(tmp_path):
    from adrd_review.eval import accuracy

    key = tmp_path / "key.csv"
    key.write_text("#,PMID,title,AI outcome\n1,a,t,included\n2,b,t,included\n3,c,t,excluded (title)\n"
                   "4,d,t,excluded (abstract)\n5,e,t,excluded (full text)\n")
    check = tmp_path / "check.csv"
    check.write_text("#,PMID,human decision (include / exclude)\n1,a,include\n2,b,exclude\n3,c,exclude\n"
                     "4,d,include\n5,e,\n")
    r = accuracy.score_check(check, key)
    assert r["scored"] == 4 and r["unsure_or_blank"] == 1
    assert r["confusion"] == {"both include": 1, "both exclude": 1, "AI include, professor exclude": 1,
                              "AI exclude, professor include (missed)": 1}
    assert r["accuracy"]["value"] == 0.5 and r["sensitivity_recall"]["value"] == 0.5
    lo, hi = r["accuracy"]["ci95"]
    assert 0 < lo < 0.5 < hi < 1
    assert accuracy.wilson(0, 0) is None and accuracy.kappa_label(0.55) == "moderate"
