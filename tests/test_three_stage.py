"""Three-stage review (titles -> abstracts -> full text), stage counts, unlabelled corpora, the AI/ML codebook."""

import json
from pathlib import Path

import pytest

from adrd_review.agents import title_screening
from adrd_review.codebook import load_codebook
from adrd_review.consensus import combine_title
from adrd_review.eval.flow import flow_counts, flow_figure, format_flow
from adrd_review.fake_model import default_responders
from adrd_review.gold import GoldRecord
from adrd_review.graph.pipeline import PipelineOptions, Reviewer, build_pipeline
from adrd_review.graph.runner import RunConfig, select_pmids, to_prediction
from adrd_review.llm import FakeLLM, ModelConfig
from adrd_review.schemas import Article, ExclusionCheck, InclusionCheck, ScreeningAnswer, Vote
from adrd_review.tools.fulltext import FullTextResult
from adrd_review.tools.pmc import FullText, Section

ROOT = Path(__file__).resolve().parents[1]
CODEBOOK = load_codebook(ROOT / "configs/codebook.yaml")
AIML = load_codebook(ROOT / "configs/codebook_aiml.yaml")

ON_TOPIC = Article(pmid="1", title="The MMSE for dementia in older adults.", abstract="We evaluated the MMSE for dementia.")
OFF_TOPIC = Article(pmid="2", title="Exercise and blood pressure in young athletes.",
                    abstract="A trial of exercise. Participants also completed the MMSE for dementia.")


def answer(inclusion, exclusion=()):
    return ScreeningAnswer(
        inclusion=[InclusionCheck(rule_id=r, status=s, quote=q) for r, s, q in inclusion],
        exclusion=[ExclusionCheck(rule_id=r, status=s, quote=q) for r, s, q in exclusion],
        rationale="r", confidence=0.9)


def tv(rid, label):
    return Vote(reviewer_id=rid, model="m", field="title", labels=[label], confidence=0.8)


# ---- the title agent and its consensus -------------------------------------------------


def test_title_panel_excludes_only_when_every_reviewer_excludes():
    assert combine_title([tv("A", "exclude"), tv("B", "exclude"), tv("C", "exclude")]).labels == ["exclude"]
    d = combine_title([tv("A", "exclude"), tv("B", "include"), tv("C", "exclude")])
    assert (d.labels, d.method) == (["include"], "any_include")  # one reviewer keeping it is enough
    assert combine_title([tv("A", "exclude")]).method == "single"


def test_title_agent_sees_only_the_title_and_quotes_must_come_from_it():
    seen = {}

    def responder(_schema, system, user):
        seen["user"], seen["system"] = user, system
        return answer([("I1", "unclear", None)])

    llm = FakeLLM(ModelConfig(provider="fake", model="x"), {ScreeningAnswer: responder})
    vote, usage = title_screening.run(OFF_TOPIC, CODEBOOK, llm)
    assert "Exercise and blood pressure" in seen["user"] and "MMSE" not in seen["user"]  # no abstract
    assert "TITLES ONLY" in seen["system"] and vote.field == "title" and usage.role == "title"
    # A not_met quoting the abstract is not evidence at the title stage -> unclear -> kept.
    from_abstract = answer([("I2", "not_met", "A trial of exercise.")])
    from adrd_review.agents import screening
    assert screening.decide(from_abstract, CODEBOOK, title_screening.title_only(OFF_TOPIC))[0] == "include"
    from_title = answer([("I2", "not_met", "Exercise and blood pressure in young athletes.")])
    assert screening.decide(from_title, CODEBOOK, title_screening.title_only(OFF_TOPIC))[0] == "exclude"


# ---- the three-stage graph -----------------------------------------------------------------


PAPER = FullText(pmid="1", pmcid="PMC1", sections=[
    Section(title="Methods", text="We recruited 412 community-dwelling adults aged 65 and over. All completed the MMSE. "
                                  "Dementia was diagnosed with DSM-IV criteria by a geriatrician blind to the MMSE."),
    Section(title="Results", text="At a cut-off of 24 the MMSE had sensitivity 0.81 and specificity 0.88 for dementia."),
])


def graph(title_stage=True, n=3, codebook=CODEBOOK, source=None):
    styles = ["keyword-rules", "keyword-strict", "keyword-lenient"]
    llms = [FakeLLM(ModelConfig(provider="fake", model=s), default_responders(codebook, s)) for s in styles[:n]]
    opts = PipelineOptions(title_stage=title_stage, eligibility=True,
                           fulltext_source=source or (lambda a: FullTextResult(PAPER, "pmc", "PMC1")))
    return build_pipeline(codebook, llms[0], llms[0], [Reviewer(c, m, m) for c, m in zip("ABC", llms)], opts)


def test_off_topic_title_stops_before_the_abstract_is_read():
    assert graph().invoke({"article": OFF_TOPIC})["title_decision"].method == "any_include"  # lenient reviewer keeps it
    s = graph(n=2).invoke({"article": OFF_TOPIC})
    assert s["title_decision"].labels == ["exclude"]
    assert not s.get("screening_votes") and "triage" not in s  # the abstract was never screened
    assert s["screening_decision"].method == "title_excluded"
    assert to_prediction(s, two_stage=True).screening == "exclude"


def test_on_topic_title_goes_through_all_three_stages():
    s = graph().invoke({"article": ON_TOPIC})
    assert s["title_decision"].labels == ["include"]
    assert s["screening_votes"] and s["eligibility_votes"]
    assert s["fulltext_info"]["source"] == "pmc"


def test_title_stage_off_changes_nothing():
    s = graph(title_stage=False).invoke({"article": OFF_TOPIC})
    assert not s.get("title_votes") and "title_decision" not in s and s["screening_votes"]


# ---- flow counts --------------------------------------------------------------------------


def test_flow_counts_each_stage_with_reasons_and_gold_losses(tmp_path):
    none = FullTextResult(None, None, "subscription", link=None)
    g = graph(n=2)
    no_text = Article(pmid="3", title="The MMSE for dementia in older adults (cohort).", abstract="We evaluated the MMSE for dementia.")
    states = [g.invoke({"article": a}) for a in (ON_TOPIC, OFF_TOPIC)]
    states.append(graph(n=2, source=lambda a: none).invoke({"article": no_text}))
    gold = {"1": GoldRecord(pmid="1", screening="include", fulltext="include"),
            "2": GoldRecord(pmid="2", screening="include", fulltext="exclude")}
    f = flow_counts(states, failures=1, codebook=CODEBOOK, gold_by=gold)
    assert f["identified"] == 4
    assert (f["title"]["screened"], f["title"]["excluded"], f["title"]["passed"]) == (3, 1, 2)
    assert f["title"]["gold_abstract_includes_lost"] == 1 and f["title"]["gold_final_includes_lost"] == 0
    assert list(f["title"]["reasons"]) == ["I1 Index test"]
    assert f["abstract"]["screened"] == 2 and f["fulltext"]["sought"] == 2
    assert f["fulltext"]["not_retrieved"] == 1 and f["fulltext"]["still_open"]["no_fulltext"] == 1
    assert f["included"]["total"] == 3 and f["included"]["waiting_for_human"] == 1 and f["included"]["failed_kept"] == 1
    assert f["included"]["gold_final_includes_kept"] == 1
    text = format_flow(f)
    assert "1. Title screening" in text and "3. Full text" in text and "INCLUDED" in text
    assert flow_figure(f, tmp_path / "flow.png").stat().st_size > 1000


def test_flow_without_title_stage_or_gold():
    s = graph(title_stage=False).invoke({"article": ON_TOPIC})
    f = flow_counts([s])
    assert "title" not in f and f["abstract"]["screened"] == 1 and "gold_final_includes_lost" not in f["abstract"]
    assert "1. Abstract screening" in format_flow(f)


# ---- unlabelled corpus and the AI/ML codebook -----------------------------------------------


def test_unlabelled_run_reads_a_pmid_list(tmp_path):
    lst = tmp_path / "pmids.txt"
    lst.write_text("111\n222\n\n111\nnot-a-pmid\n")
    cfg = RunConfig(name="x", gold=None, pmids_file=str(lst))
    assert select_pmids(cfg, "all", []) == ["111", "222"]
    with pytest.raises(ValueError, match="gold file .* or a pmids_file"):
        select_pmids(RunConfig(name="y", gold=None), "all", [])


def test_human_gold_refused_without_gold():
    from adrd_review.graph.runner import build_graph

    fake = {"model": {"provider": "fake", "model": "keyword-rules"}}
    cfg = RunConfig(name="x", gold=None, human="gold", roles={"triage": fake, "extraction": fake, "screening": fake, "annotation": fake})
    with pytest.raises(ValueError, match="needs a gold standard"):
        build_graph(cfg, CODEBOOK, [])


def test_aiml_codebook_has_the_professors_criteria_and_names_the_real_topic():
    assert AIML.version == "0.3.0" and "machine learning" in AIML.topic
    assert [r.id for r in AIML.screening.inclusion] == ["I1", "I2", "I3", "I4"]
    assert all(r.keywords for r in AIML.screening.inclusion)
    rules = {r.id: " ".join(r.text.split()) for r in AIML.rules}
    assert "MCI alone" in rules["I1"] or "MCI) counts ONLY together" in rules["I1"]  # professor, 2026-10-07
    assert "Logistic regression counts" in rules["I2"] and "60 or over" in rules["I4"]
    assert "Full papers in conference proceedings are NOT excluded" in rules["E2"]
    assert [r.id for r in AIML.rules if "PROPOSED" in r.source] == ["I4", "E3", "E4"]
    from adrd_review.agents import screening
    assert "scoping review of AI and machine learning" in screening.system_prompt(AIML)
    assert "keywords" not in screening.system_prompt(AIML).lower()  # stand-in hints never reach a real model
    # the placeholder codebook's prompts are unchanged (cached answers stay valid)
    assert "scoping review of dementia screening research" in screening.system_prompt(CODEBOOK)


def test_fake_model_follows_the_aiml_codebook_through_three_stages():
    ml = Article(pmid="7", title="Deep learning on MRI to detect Alzheimer's disease in older adults.",
                 abstract="We trained a convolutional neural network on ADNI MRI scans of 800 older adults to predict dementia.")
    review = Article(pmid="8", title="Machine learning for dementia: a systematic review.", abstract="We reviewed 40 studies.")
    animal = Article(pmid="9", title="Gut microbes in young mice.", abstract="Mice were fed a diet.")
    g = graph(codebook=AIML, source=lambda a: FullTextResult(None, None, "none", link=None))
    s = g.invoke({"article": ml})
    assert s["title_decision"].labels == ["include"] and s["screening_decision"].labels == ["include"]
    assert "ADNI" in s["annotation_decisions"]["dataset_source"].labels
    assert g.invoke({"article": review})["screening_decision"].method == "title_excluded"  # E1 from the title
    assert g.invoke({"article": animal})["screening_decision"].method == "title_excluded"


@pytest.mark.skipif(not (ROOT / "data/search/prof_query_pmids.txt").exists(), reason="needs the professor's search results")
def test_preflight_estimates_the_unlabelled_search_run(monkeypatch):
    from adrd_review.preflight import check

    monkeypatch.chdir(ROOT)
    rep = check("configs/runs/PROF_search_fake.yaml", calls=False, limit=20)
    assert not rep.failed
    assert any("PROPOSED" in t for _, t in rep.lines) and any("unlabelled" in t for _, t in rep.lines)
    assert not any("extraction" in t or "annotation" in t for _, t in rep.lines)  # counts-only run
    assert json.loads(json.dumps(rep.lines))  # plain data


def test_annotate_off_stops_after_screening():
    llm = FakeLLM(ModelConfig(provider="fake", model="keyword-rules"), default_responders(CODEBOOK, "keyword-rules"))
    opts = PipelineOptions(title_stage=True, eligibility=True, annotate=False,
                           fulltext_source=lambda a: FullTextResult(PAPER, "pmc", "PMC1"))
    s = build_pipeline(CODEBOOK, llm, llm, [Reviewer("A", llm, llm)], opts).invoke({"article": ON_TOPIC})
    assert s["eligibility_decision"].labels == ["include"]
    assert "extraction" not in s and not s.get("annotation_votes")
    assert {u.role for u in s["usage"]} == {"title", "triage", "screening", "eligibility"}
    assert flow_counts([s])["included"]["confirmed"] == 1
    # one stage only (no full text) also stops before extraction
    opts1 = PipelineOptions(annotate=False)
    s1 = build_pipeline(CODEBOOK, llm, llm, [Reviewer("A", llm, llm)], opts1).invoke({"article": ON_TOPIC})
    assert s1["screening_decision"].labels == ["include"] and "extraction" not in s1


def test_missing_fulltext_not_retrieved_is_left_out_and_counted():
    llm = FakeLLM(ModelConfig(provider="fake", model="keyword-rules"), default_responders(CODEBOOK, "keyword-rules"))
    none = FullTextResult(None, None, "subscription", link=None)

    def run(mode, source):
        opts = PipelineOptions(eligibility=True, annotate=False, missing_fulltext=mode, verify=True, fulltext_source=source)
        return build_pipeline(CODEBOOK, llm, llm, [Reviewer("A", llm, llm)], opts).invoke({"article": ON_TOPIC})

    dropped = run("not_retrieved", lambda a: none)
    assert (dropped["eligibility_decision"].labels, dropped["eligibility_decision"].method) == (["exclude"], "not_retrieved")
    assert not dropped.get("escalations")  # nobody is asked about it
    kept = run("human", lambda a: none)
    assert kept["eligibility_decision"].method == "no_fulltext"  # default: kept for a person
    found = run("not_retrieved", lambda a: FullTextResult(PAPER, "pmc", "PMC1"))
    f = flow_counts([dropped, kept, found])["fulltext"]
    assert (f["sought"], f["retrieved"], f["not_retrieved"], f["not_retrieved_left_out"]) == (3, 1, 2, 1)
    assert f["excluded"] == 0 and f["passed"] == 2 and f["still_open"]["no_fulltext"] == 1
    assert flow_counts([dropped, found])["included"]["total"] == 1
    assert "(left out)" in format_flow(flow_counts([dropped, found]))


def test_temperature_can_be_left_out_for_models_that_reject_it():
    from adrd_review.llm import make_chat_model

    m = make_chat_model(ModelConfig(provider="anthropic", model="claude-haiku-5-5", temperature=None, extra={"api_key": "x"}))
    assert getattr(m, "temperature", None) is None
    m = make_chat_model(ModelConfig(provider="anthropic", model="claude-haiku-5-5", extra={"api_key": "x"}))
    assert m.temperature == 0.0


def test_sample_is_fixed_and_keeps_list_order(tmp_path, monkeypatch):
    import random

    from adrd_review.graph import runner

    pmids = [str(100000 + i) for i in range(50)]
    chosen = set(random.Random(runner.SAMPLE_SEED).sample(pmids, 10))
    assert chosen == set(random.Random(runner.SAMPLE_SEED).sample(pmids, 10))  # same pilot every time
    assert [p for p in pmids if p in chosen] == sorted(chosen, key=pmids.index)


def test_out_of_credit_stops_at_once_and_skips_remaining_articles():
    from adrd_review.graph.runner import run_articles
    from adrd_review.llm import OUT_OF_CREDIT, LangChainLLM, OutOfCredit, ResponseCache

    calls = []

    class Broke:
        def with_structured_output(self, *a, **k):
            return self

        def invoke(self, _messages):
            calls.append(1)
            raise RuntimeError("Error code: 400 - Your credit balance is too low to access the Anthropic API.")

    OUT_OF_CREDIT.clear()
    try:
        llm = LangChainLLM(ModelConfig(provider="anthropic", model="m", temperature=None), chat_model=Broke(),
                           cache=ResponseCache(None))
        with pytest.raises(OutOfCredit):
            title_screening.run(ON_TOPIC, CODEBOOK, llm)
        assert len(calls) == 1  # not retried
        g = build_pipeline(CODEBOOK, llm, llm, [Reviewer("A", llm, llm)], PipelineOptions())
        _, failures, _ = run_articles(g, [ON_TOPIC, OFF_TOPIC], concurrency=1)
        assert len(failures) == 2 and all("OutOfCredit" in f["error"] for f in failures)
        assert len(calls) == 1  # the run started nothing new
    finally:
        OUT_OF_CREDIT.clear()


def test_anthropic_prompt_cache_marks_the_instructions_and_prices_cached_tokens():
    from types import SimpleNamespace

    from adrd_review.llm import LangChainLLM, ResponseCache

    seen = {}

    class Chat:
        def with_structured_output(self, *a, **k):
            return self

        def invoke(self, messages):
            seen["system"] = messages[0].content
            raw = SimpleNamespace(usage_metadata={"input_tokens": 3000, "output_tokens": 100,
                                                  "input_token_details": {"cache_read": 2500, "cache_creation": 0}})
            return {"parsed": answer([("I1", "unclear", None)]), "raw": raw}

    cfg = ModelConfig(provider="anthropic", model="claude-sonnet-5-5", temperature=None, prompt_cache=True)
    r = LangChainLLM(cfg, chat_model=Chat(), cache=ResponseCache(None)).structured(ScreeningAnswer, "RULES", "PAPER", role="title")
    assert seen["system"] == [{"type": "text", "text": "RULES", "cache_control": {"type": "ephemeral"}}]
    # 500 fresh tokens at $2/M + 2,500 cached at $0.10/M + 100 out at $10/M (configs/prices.yaml)
    assert r.usage.prompt_cache_read_tokens == 2500
    assert r.usage.cost_usd == pytest.approx((500 * 2 + 2500 * 0.10 + 100 * 10) / 1e6)
    plain = LangChainLLM(cfg.model_copy(update={"prompt_cache": False}), chat_model=Chat(), cache=ResponseCache(None))
    plain.structured(ScreeningAnswer, "RULES", "PAPER", role="title")
    assert seen["system"] == "RULES"


def test_stage_files_match_the_flow_and_hide_the_ai_columns(tmp_path):
    from openpyxl import load_workbook

    from adrd_review.eval import stages
    from adrd_review.graph.runner import _dump

    on, off = ON_TOPIC, OFF_TOPIC.model_copy(update={"pmid": "2"})
    nofull = Article(pmid="3", title="The MMSE for dementia in older adults, again.", abstract="We evaluated the MMSE for dementia.")
    source = lambda a: FullTextResult(PAPER, "pmc", "PMC1") if a.pmid != "3" else FullTextResult(None, None, "subscription", link=None)
    llms = [FakeLLM(ModelConfig(provider="fake", model=s), default_responders(CODEBOOK, s)) for s in ["keyword-rules", "keyword-strict"]]
    opts = PipelineOptions(title_stage=True, eligibility=True, annotate=False, missing_fulltext="not_retrieved", fulltext_source=source)
    g = build_pipeline(CODEBOOK, llms[0], llms[0], [Reviewer(c, m, m) for c, m in zip("AB", llms)], opts)
    states = [g.invoke({"article": a}) for a in (on, off, nofull)]
    (tmp_path / "decisions.jsonl").write_text(
        "\n".join(json.dumps(_dump({k: v for k, v in s.items() if k != "usage"} | {"pmid": s["article"].pmid})) for s in states))
    (tmp_path / "failures.json").write_text("[]")

    out = stages.export(tmp_path, n=1, seed=1, codebook=CODEBOOK)
    rows = stages.collect(tmp_path, CODEBOOK)
    flow = flow_counts(states, 0, CODEBOOK)
    assert len(rows["title"]) == flow["title"]["screened"] == 3
    assert [r["pmid"] for r in rows["title"] if r["ai_outcome"] == "excluded"] == ["2"]
    assert rows["title"][1]["main reason"].startswith(next(iter(flow["title"]["reasons"]))[:2])
    assert rows["title"][0]["abstract"] == "" and rows["abstract"][0]["abstract"]  # titles: what the reviewers saw
    assert {r["pmid"]: r["ai_outcome"] for r in rows["fulltext"]} == {"1": "passed", "3": "not retrieved"}
    assert [r["pmid"] for r in rows["included"]] == ["1"] == [str(flow["included"]["total"])]
    assert all(r["human_decision"] == "" for rs in rows.values() for r in rs)

    wb = load_workbook(out / "stages.xlsx")
    assert wb.sheetnames == ["random sample", "1 titles", "2 abstracts", "3 full text", "4 included", "how to check"]
    ws = wb["2 abstracts"]
    header = [c.value for c in ws[1]]
    ai = ws.cell(row=1, column=header.index("ai_outcome") + 1).column_letter
    assert ws.column_dimensions[ai].hidden  # blind: the AI's decision is hidden until unhidden
    assert header.index("human_decision") < header.index("ai_outcome")
    picked = stages.sample(rows, n=1, seed=1)
    assert picked == stages.sample(rows, n=1, seed=1)  # the same seed draws the same papers
    assert {(r["stage"], r["ai_outcome"]) for r in picked} == {(s, r["ai_outcome"]) for s, rs in rows.items() for r in rs
                                                               if r["ai_outcome"] != "not retrieved"}  # free full text only
    import csv
    assert len(list(csv.DictReader((out / "1_title.csv").open()))) == 3


def test_an_article_that_fails_part_way_keeps_its_finished_stages_in_the_stage_files(tmp_path):
    from adrd_review.eval import stages
    from adrd_review.graph.runner import _dump, _title_only, run_articles

    def broken(article):
        raise RuntimeError("daily limit used up")

    llm = FakeLLM(ModelConfig(provider="fake", model="keyword-rules"), default_responders(CODEBOOK, "keyword-rules"))
    opts = PipelineOptions(title_stage=True, eligibility=True, annotate=False, fulltext_source=broken)
    g = build_pipeline(CODEBOOK, llm, llm, [Reviewer("A", llm, llm)], opts)
    states, failures, _ = run_articles(g, [ON_TOPIC], concurrency=1)
    (fail,) = failures
    assert not states and "daily limit" in fail["error"]
    partial = fail["partial"]  # title and abstract were decided before the full text failed
    assert partial["title_decision"].labels == ["include"] and partial["screening_decision"].labels == ["include"]

    # a deferred article (its abstract batch failed) keeps its title decision from the batch votes
    a = Reviewer("A", llm, llm)
    later = Article(pmid="5", title="Deep learning for dementia.", abstract="x")
    initial = {"5": {"prepared": {"title": {"A": tv("A", "include")}}, "usage": []}}
    assert _title_only(later, initial, [a])["partial"]["title_decision"].labels == ["include"]
    assert _title_only(later, {"5": {"prepared": {}}}, [a]) == {}  # title stage not finished: nothing to show

    open_rows = [_dump(partial) | {"pmid": "1", "open": fail["error"]},
                 _dump(_title_only(later, initial, [a])["partial"]) | {"pmid": "5", "open": "Deferred: abstract batch"}]
    (tmp_path / "in_progress.jsonl").write_text("\n".join(json.dumps(r) for r in open_rows))
    (tmp_path / "decisions.jsonl").write_text("")
    rows = stages.collect(tmp_path, CODEBOOK)
    assert {r["pmid"]: r["ai_outcome"] for r in rows["title"]} == {"1": "passed", "5": "passed"}
    assert {r["pmid"]: r["ai_outcome"] for r in rows["abstract"]} == {"1": "passed", "5": "still open"}
    assert {r["pmid"]: r["ai_outcome"] for r in rows["fulltext"]} == {"1": "still open"} and not rows["included"]
    assert rows["fulltext"][0]["how decided"].startswith("open: RuntimeError")
    assert all(r["ai_outcome"] != "still open" for r in stages.sample(rows, n=5))  # nothing to check yet


def test_second_look_rereads_only_the_papers_left_for_a_person(tmp_path):
    from adrd_review import second_look
    from adrd_review.llm import LLMError

    paper = Article(pmid="7", title="Deep learning on MRI to detect Alzheimer's disease in older adults.",
                    abstract="We trained a CNN on ADNI MRI scans to predict dementia.",
                    fulltext="We trained a convolutional neural network on ADNI MRI scans of 800 older adults aged 65 "
                             "and over to detect Alzheimer's disease and predict dementia.")

    def state(pmid, decision, votes):
        return {"pmid": pmid, "article": paper.model_copy(update={"pmid": pmid}).model_dump(),
                "title_decision": {"labels": ["include"]}, "screening_decision": {"labels": ["include"]},
                "fulltext_info": {"stage": 2, "used": True}, "eligibility_decision": {"labels": [decision]},
                "eligibility_votes": votes}

    unclear = [{"reviewer_id": "A", "labels": ["unclear"], "rule_ids": ["I4"]},
               {"reviewer_id": "B", "labels": ["include"], "rule_ids": ["I1", "I2", "I3", "I4"]}]
    rows = [state("7", "unclear", unclear), state("8", "include", unclear[1:])]
    (tmp_path / "decisions.jsonl").write_text("\n".join(json.dumps(r) for r in rows))
    (tmp_path / "run_meta.json").write_text(json.dumps({"config": {"codebook": str(ROOT / "configs/codebook_aiml.yaml")}}))

    assert second_look.open_rules(rows[0], AIML) == ["I4"]
    llm = FakeLLM(ModelConfig(provider="fake", model="keyword-rules"), default_responders(AIML, "keyword-rules"))
    (row,) = second_look.run(tmp_path, llm=llm)  # only the paper left for a person, not the included one
    assert row["pmid"] == "7" and row["open rules"] == "I4"
    assert row["second look"] in second_look.OUTCOME.values() and row["model"] == "fake:keyword-rules"
    assert (tmp_path / "second_look.csv").exists()

    class Down:
        config = ModelConfig(provider="fake", model="down")

        def structured(self, *a, **k):
            raise LLMError("down")

    (row,) = second_look.run(tmp_path, llm=Down())
    assert row["second look"] == "for human review" and "failed" in row["reasoning"]  # never guessed


def test_age_check_applies_65_from_quoted_ages_only():
    from adrd_review import age_check as ac

    paper = Article(pmid="9", title="t", abstract="a",
                    fulltext="Participants aged 65 years or older were enrolled. The mean age was 72.4 years. "
                             "In the second cohort the mean age was 61.5 years.")
    fact = ac.AgeFact
    assert ac.decide(ac.AgeAnswer(min_age=fact(value=65, quote="aged 65 years or older")), paper)[0] == "met"
    assert ac.decide(ac.AgeAnswer(mean_age=fact(value=72.4, quote="The mean age was 72.4 years")), paper)[0] == "met"
    assert ac.decide(ac.AgeAnswer(mean_age=fact(value=61.5, quote="the mean age was 61.5 years")), paper)[0] == "not_met"
    # a number whose quote is not in the paper, or does not contain the number, counts for nothing
    assert ac.decide(ac.AgeAnswer(mean_age=fact(value=70, quote="The mean age was 70 years")), paper)[0] == "unclear"
    assert ac.decide(ac.AgeAnswer(mean_age=fact(value=80, quote="The mean age was 72.4 years")), paper)[0] == "unclear"
    assert ac.decide(ac.AgeAnswer(), paper)[0] == "unclear"
    assert ac.open_only_on_age({"eligibility_votes": [{"labels": ["unclear"], "rule_ids": ["I4"]},
                                                      {"labels": ["include"], "rule_ids": ["I1", "I2"]}]})
    assert not ac.open_only_on_age({"eligibility_votes": [{"labels": ["unclear"], "rule_ids": ["I2", "I4"]}]})


def test_final_results_summary_counts_each_paper_once_and_keeps_paid_access_aside():
    from adrd_review import final_results as fr

    final = [{"stage": "title", "outcome": "excluded", "reason": "E1 Review", "decided by": "panel"},
             {"stage": "abstract", "outcome": "excluded", "reason": "I4 Older adults (65 or over)",
              "decided by": "age check on the abstract"},
             {"stage": "full text", "outcome": "included", "reason": "all rules met", "decided by": "panel"},
             {"stage": "full text", "outcome": "for human review", "reason": "age", "decided by": "x"},
             {"stage": "full text", "outcome": fr.PAID, "reason": "no free full text", "decided by": ""},
             {"stage": "full text", "outcome": fr.HAND, "reason": "blocked", "decided by": ""}]
    s = fr.summarize(final)
    assert s["title"] == {"screened": 6, "excluded": 1, "passed": 5}
    assert s["abstract"]["passed"] == 4 and s["abstract"]["excluded_by_age_check"] == 1
    assert s["fulltext_free"]["assessed"] == 2 and s["fulltext_free"]["included"] == 1  # paid / by-hand not counted
    assert s["side_note"] == {"free_copy_to_save_by_hand": 1, "paid_access_not_assessed": 1}
    assert fr._failing_rule("I1 met · I3 NOT MET · E1 APPLIES", {"I1": "I1 x", "I3": "I3 y", "E1": "E1 z"}) == "I3 y"


def test_third_look_needs_65_plus_and_every_other_rule_met():
    from adrd_review import age_check as ac
    from adrd_review import third_look as tl
    from adrd_review.llm import LLMResult
    from adrd_review.schemas import Usage

    text = ("We trained a deep learning model on MRI to detect Alzheimer's disease. Participants were 400 adults "
            "with a mean age of 74.2 years. This is an original study of human participants.")
    paper = Article(pmid="5", title="Deep learning MRI for Alzheimer's disease", abstract="a", fulltext=text)
    state = {"pmid": "5", "article": paper.model_dump()}

    def stub(age, marks):
        class Stub:
            config = ModelConfig(provider="fake", model="stub")

            def structured(self, schema, system, user, role):
                if schema is ac.AgeAnswer:
                    return LLMResult(age, Usage(role=role, model="fake:stub"))
                ans = answer([(r, s, q) for r, s, q in marks if r.startswith("I")],
                             [(r, s, q) for r, s, q in marks if r.startswith("E")])
                return LLMResult(ans, Usage(role=role, model="fake:stub"))
        return Stub()

    met = [("I1", "met", "detect Alzheimer's disease"), ("I2", "met", "deep learning model"),
           ("I3", "met", "to detect Alzheimer's disease"), ("I4", "unclear", None),
           ("E1", "does_not_apply", None), ("E2", "does_not_apply", None), ("E3", "does_not_apply", None),
           ("E4", "does_not_apply", None)]
    age74 = ac.AgeAnswer(mean_age=ac.AgeFact(value=74.2, quote="a mean age of 74.2 years"))
    assert tl.look(state, {}, "age", AIML, stub(age74, met))["third look"] == "included"
    assert tl.look(state, {}, "age", AIML, stub(ac.AgeAnswer(), met))["third look"] == "for human review"
    # an earlier verified age check counts when this model finds no age
    assert tl.look(state, {"I4 at 65": "met", "reason": "mean 70"}, "age", AIML, stub(ac.AgeAnswer(), met))["third look"] == "included"
    not_ai = [m if m[0] != "I2" else ("I2", "not_met", "This is an original study of human participants") for m in met]
    assert tl.look(state, {}, "age", AIML, stub(age74, not_ai))["third look"] == "excluded"
